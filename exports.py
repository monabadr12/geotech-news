"""
Excel workbooks for the archive: one per issue window, plus one for the whole
archive. Built with openpyxl, no macros; works in Excel 2016+ / Microsoft 365
and LibreOffice.

Sheets (in tab order):
  Start here  what each sheet is for, how to use the selectors, links to
              every sheet and to each theme section
  Dashboard   dropdown selectors (theme, suggested use, minimum relevance and,
              in the whole-archive workbook, issue window). Every figure,
              table, chart and the Top 10 list is a formula over the Stories
              table, so they all follow the selectors and the Status column.
  Stories     every story as an Excel table, with planning columns (Status,
              Assigned to, Editor notes) the editor fills in, and where it
              happened (country, place) and the developing story it belongs to
  By theme    read-only view, one collapsible section per theme
  Research    new papers from the journals (when there are any)
  Events      the events calendar: open call-for-papers deadlines and the
              upcoming events (when the calendar has any)
  Sources     the outlets that supplied stories
  Lists       (hidden) the values offered by the dropdowns, plus helper cells

Every piece of text from the news is written as a plain string (never as a
formula) with control characters removed.
"""
import datetime as dt
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.descriptors import Typed
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import Outline
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
NAVY, NAVY2, ACCENT, INK, MUTED, LINE = "16324F", "23466B", "C9A227", "1B2430", "5B6573", "D9DEE3"
LINK, PAPER, INPUT_FILL, EDIT_FILL, WHITE = "1D5A8A", "F4F6F9", "FFF4CC", "FFFBEF", "FFFFFF"
NOT_RATED = "Not rated"
STANDARD_USES = ["Feature lead", "News brief", "Academia-to-Industry topic", "Regional insight", "Watch only"]
USE_COLORS = {"Feature lead": "D38B00", "News brief": "2563D6", "Academia-to-Industry topic": "6B4FD6",
              "Regional insight": "1F9D63", "Watch only": "8A94A3", NOT_RATED: "B8C0CA", "": "B8C0CA"}
REL_COLORS = {5: "1E7B4F", 4: "4DA06E", 3: "D9A21B", 2: "A3ADB8", 1: "C9CED6", 0: "B8C0CA"}
REL_LABELS = {5: "5 / 5", 4: "4 / 5", 3: "3 / 5", 2: "2 / 5", 1: "1 / 5", 0: "Keyword-ranked"}
REL_MEANING = {5: "Rich geotechnical lesson or a major ground-related event",
               4: "Strong ground-engineering content",
               3: "Clear ground-engineering element, worth a brief",
               2: "Weak link to geotechnics",
               1: "Not really geotechnical",
               0: "Shown as –. Not rated by AI: ranked on keyword matches only"}
# (status, text colour, cell fill, chart colour, meaning)
STATUSES = [("To review", "4A5563", "ECEFF3", "A3ADB8", "New from the scan, not looked at yet (every story starts here)"),
            ("Shortlisted", "1D4FB8", "DCE7FB", "2563D6", "Worth considering for the issue"),
            ("Commissioned", "8A5300", "FBE6C2", "D38B00", "A writer or contributor has been asked to cover it"),
            ("Published", "17663F", "D3EDDF", "1E7B4F", "Used in a published issue or newsletter"),
            ("Not used", "7A8491", "F1F2F4", "D5DAE0", "Reviewed and set aside")]
DEFAULT_STATUS = STATUSES[0][0]
KPI_COLORS = [NAVY, "0B97A5", "1E7B4F", "D38B00", "6B4FD6", "2563D6"]
TAB_COLORS = {"Start here": ACCENT, "Dashboard": NAVY, "Stories": "2563D6", "By theme": "1F9D63",
              "Research": "6B4FD6", "Events": "D38B00", "Sources": "0B97A5", "Lists": "B8C0CA"}
# the events calendar's kinds of event (events.KINDS): text colour on the Events sheet
KIND_COLORS = {"ISSMGE conference": "7D5A0C", "ISSMGE-supported": "0A6D78", "Other conference": "36506C",
               "Course / webinar": "5137B0"}
ALL_THEMES, ALL_USES, ANY_REL, ALL_WINDOWS = "All themes", "All uses", "Any", "All windows"
THIN = Side(style="thin", color=LINE)
MAX_TEXT = 32000          # Excel holds at most 32,767 characters in a cell
MAX_LINKS = 60000         # Excel allows 65,530 hyperlinks per sheet

# Dashboard selector cells (merged blocks: the anchors hold the values)
SEL_THEME, SEL_USE, SEL_REL, SEL_WIN = "$B$5", "$G$5", "$K$5", "$N$5"
DASH = "Dashboard!"

# Stories sheet columns: key, header, width
STORY_COLS = [   # the printed part runs from No. to Digest date
    ("no", "No.", 6), ("headline", "Headline (English)", 46), ("theme", "Theme", 21), ("rel", "Relevance", 10),
    ("use", "Suggested use", 16), ("status", "Status", 14), ("assigned", "Assigned to", 14),
    ("notes", "Editor notes", 28), ("angle", "Magazine angle", 40), ("source", "Source", 19),
    ("digest", "Digest date", 12), ("original", "Original headline", 38), ("summary", "Summary", 40),
    ("country", "Country", 16), ("place", "Place", 22), ("storyline", "Developing story", 36),
    ("edition", "Google News edition", 15), ("published", "Published", 12), ("window", "Issue window", 11),
    ("keywords", "Keywords", 26), ("score", "Keyword score", 9), ("url", "Article URL", 30),
    ("insel", "In Dashboard selection", 12), ("key", "Top-10 key (helper)", 12),
]
PRINT_TO = "digest"
SC = {key: get_column_letter(i) for i, (key, _, _) in enumerate(STORY_COLS, 1)}
SCI = {key: i for i, (key, _, _) in enumerate(STORY_COLS, 1)}
PLANNING = ("status", "assigned", "notes")
WRAPPED = ("headline", "theme", "use", "notes", "angle", "source", "original", "summary", "keywords", "place", "storyline")
STORY_HEADER, STORY_FIRST = 3, 4


class F(str):
    """A formula to be written as a formula (all other text is forced to be a string)."""


def q(text):
    """A text literal inside a formula."""
    return '"' + str(text).replace('"', '""') + '"'


# ------------------------------------------------------------------ helpers
def hexc(color):
    return str(color).lstrip("#").upper()


def tint(color, amount=0.82):
    c = hexc(color)
    rgb = [int(c[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(v + (255 - v) * amount):02X}" for v in rgb)


def fill(color):
    return PatternFill("solid", start_color=hexc(color), end_color=hexc(color))


def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=hexc(color), italic=italic, underline=underline)


def clean(text, limit=MAX_TEXT):
    """Plain text that openpyxl can store: control characters removed, length capped."""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(text))
    return text if len(text) <= limit else text[:limit - 1] + "…"


def shorten(text, limit):
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def as_date(iso):
    try:
        return dt.date.fromisoformat(str(iso)[:10]) if iso else None
    except ValueError:
        return None


def safe_url(url):
    """The article link if it is a web address Excel can hold (at most ~2,000 characters), else ''."""
    url = ILLEGAL_CHARACTERS_RE.sub("", str(url or "")).strip()
    return url if len(url) <= 2000 and re.match(r"(?i)^(https?://|mailto:)\S+$", url) else ""


def put(ws, row, col, value=None, *, f=None, fl=None, al=None, fmt=None, border=None):
    """Write one cell. Text is always stored as text, so a headline such as '=1+1' stays a headline."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, F):
        cell.value = str(value)
    elif isinstance(value, str):
        cell.value = clean(value)
        cell.data_type = "s"
    elif value is not None:
        cell.value = value
    cell.font = f or font()
    if fl:
        cell.fill = fill(fl)
    if al:
        cell.alignment = al
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def span(ws, row, c1, c2, value=None, *, f=None, fl=None, al=None, fmt=None, border=None, r2=None):
    """Write value into a merged block (row..r2, c1..c2), styling every cell of the block."""
    r2 = r2 or row
    if c2 > c1 or r2 > row:
        ws.merge_cells(start_row=row, start_column=c1, end_row=r2, end_column=c2)
    for r in range(row, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if fl:
                cell.fill = fill(fl)
            if border:
                cell.border = border
    return put(ws, row, c1, value, f=f, fl=fl, al=al, fmt=fmt, border=border)


def box(ws, r1, c1, r2, c2, side):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = Border(left=side if c == c1 else None, right=side if c == c2 else None,
                                                     top=side if r == r1 else None, bottom=side if r == r2 else None)


def link_to(cell, sheet, ref="A1"):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{ref}", display=str(cell.value or ""))


def nav(ws, row, col, items):
    """A row of internal links such as '← Dashboard'."""
    for i, (text, sheet, ref) in enumerate(items):
        c = put(ws, row, col + i, text, f=font(9, True, LINK, underline="single"),
                al=Alignment(horizontal="left", vertical="center"))
        link_to(c, sheet, ref)


def section(ws, row, c1, c2, text, height=22):
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).border = Border(bottom=Side(style="medium", color=ACCENT))
    put(ws, row, c1, text, f=font(12, True, NAVY), al=Alignment(vertical="bottom"),
        border=Border(bottom=Side(style="medium", color=ACCENT)))
    ws.row_dimensions[row].height = height


def head(ws, row, cells, height=20, color=NAVY, text=WHITE):
    """cells: [(c1, c2, label)] – a header band, with merged labels."""
    for c1, c2, label in cells:
        span(ws, row, c1, c2, label, f=font(9, True, text), fl=color,
             al=Alignment(horizontal="left" if c2 > c1 else "center", vertical="center", wrap_text=True, indent=1 if c2 > c1 else 0))
    ws.row_dimensions[row].height = height


def page_setup(ws, title_rows=None):
    """A4 landscape, one page wide, repeated header rows, sheet name and page number in the footer."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.55, header=0.25, footer=0.25)
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.left.text, ws.oddFooter.left.size, ws.oddFooter.left.font = "&A", 8, FONT
    ws.oddFooter.right.text, ws.oddFooter.right.size, ws.oddFooter.right.font = "Page &P of &N", 8, FONT


def rich(size=900, color=INK, bold=False):
    cp = CharacterProperties(sz=size, b=bold, solidFill=hexc(color), latin=DrawingFont(typeface=FONT))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text):
    cp = CharacterProperties(sz=1000, b=True, solidFill=NAVY, latin=DrawingFont(typeface=FONT))
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


class Labels(DataLabelList):
    """Data labels whose number format is not linked to the cells, so '#,##0;;;' hides zero labels."""
    tagname = "dLbls"
    numFmt = Typed(expected_type=NumFmt, allow_none=True)
    __elements__ = DataLabelList.__elements__
    __nested__ = tuple(n for n in DataLabelList.__nested__ if n != "numFmt")
    __attrs__ = DataLabelList.__attrs__


def bar_chart(ws, anchor, cats, series, *, title, width, height, horizontal=False, colors=None, series_colors=None,
              legend=False, overlap=None, gap=55):
    """series: [Reference incl. header cell]; colors: one per category (single series)."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = chart_title(title)
    for ref in series:
        ch.add_data(ref, titles_from_data=True)
    ch.set_categories(cats)
    ch.gapWidth = gap
    if overlap is not None:
        ch.overlap = overlap
    ch.x_axis.delete, ch.y_axis.delete = False, True
    ch.x_axis.majorGridlines = ch.y_axis.majorGridlines = None
    ch.x_axis.txPr = rich(850, MUTED)
    ch.x_axis.majorTickMark = ch.x_axis.minorTickMark = "none"
    ch.y_axis.scaling.min = 0
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
    ch.legend = None
    if legend:
        ch.legend = Legend(legendPos="t", overlay=False, txPr=rich(850, MUTED))
    for i, s in enumerate(ch.series):
        color = (series_colors or [ACCENT])[i % len(series_colors or [ACCENT])]
        s.graphicalProperties.solidFill = hexc(color)
        s.graphicalProperties.line.solidFill = hexc(color)
        s.invertIfNegative = False
        s.dLbls = Labels(showVal=True, showPercent=False, showSerName=False, showCatName=False,
                         showLegendKey=False, showLeaderLines=False, numFmt=NumFmt(formatCode="#,##0;;;", sourceLinked=False),
                         dLblPos="outEnd", txPr=rich(850, INK, True))
        if legend and i == 0 and len(ch.series) > 1:
            s.dLbls = None       # the background ("all stories") series stays unlabelled
    if colors and ch.series:
        for i, color in enumerate(colors):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = hexc(color)
            pt.graphicalProperties.line.solidFill = hexc(color)
            ch.series[0].dPt.append(pt)
    ch.width, ch.height = width, height
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(x=0.02, y=0.16 if legend else 0.13, w=0.96,
                                                           h=0.72 if legend else 0.8))
    ws.add_chart(ch, anchor)
    return ch


def lines_needed(text, width):
    if not text:
        return 1
    per_line = max(int(width * 1.05), 8)      # bold Arial 10 fits about one character per width unit
    return sum(max(1, -(-len(part) // per_line)) for part in str(text).split("\n"))


# ------------------------------------------------------------------ data
def _int(value, default=0):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _story(s, no):
    title = str(s.get("title") or "")
    rel = max(0, min(5, _int(s.get("relevance"))))
    summary = " ".join(str(s.get("summary") or "").split())
    return {
        "no": no, "headline": str(s.get("title_en") or title or "(no headline)"),
        "original": title if s.get("title_en") else "", "theme": str(s.get("theme") or "General"), "rel": rel,
        "use": str(s.get("use") or "") or NOT_RATED, "angle": str(s.get("angle") or ""),
        "summary": shorten(summary, 400) if summary else "", "source": str(s.get("source") or "Unknown source"),
        "edition": str(s.get("_region") or s.get("region") or ""), "published": as_date(s.get("published")),
        "digest": as_date(s.get("_digest")), "digest_iso": str(s.get("_digest") or ""),
        "window": str(s.get("_window") or ""), "keywords": ", ".join(str(k) for k in (s.get("keywords") or [])),
        "score": _int(s.get("score")), "rank": _int(s.get("rank")), "url": safe_url(s.get("link")),
        # where it happened, as Claude read it (_country: the English name the site gives the code)
        "country": str(s.get("_country") or s.get("country") or ""), "place": " ".join(str(s.get("place") or "").split()),
        "storyline": _storyline(s.get("_storyline")),
    }


def _storyline(sl):
    """'Moniquirá landslide, Colombia (Oct 2026) · Investigation & inquiry · 4 reports' (storylines.annotate)."""
    if not isinstance(sl, dict) or not sl.get("title"):
        return ""
    return " · ".join(x for x in (str(sl["title"]), str(sl.get("stage") or ""),
                                  f"{_int(sl.get('size'))} reports" if _int(sl.get("size")) else "") if x)


def _paper(p, no):
    return {
        "no": no, "title": str(p.get("title") or "(no title)"), "theme": str(p.get("theme") or "General"),
        "rel": max(0, min(5, _int(p.get("relevance")))), "angle": str(p.get("angle") or ""),
        "journal": str(p.get("journal") or ""), "authors": str(p.get("authors") or ""),
        "published": as_date(p.get("published")), "digest": as_date(p.get("_digest")),
        "abstract": shorten(p.get("abstract") or "", 800) if p.get("abstract") else "",
        "keywords": ", ".join(str(k) for k in (p.get("keywords") or [])), "url": safe_url(p.get("link")),
    }


def _timeline_key(label):
    """('date', date) for a daily-digest label such as '24 Oct 2026', ('window', '2026-Q4') for 'Q4 2026'."""
    text = " ".join(str(label).split())
    for fmt in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%A %d %B %Y"):
        try:
            return "date", dt.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    m = re.fullmatch(r"Q([1-4]) (\d{4})", text) or re.fullmatch(r"(\d{4})-Q([1-4])", text)
    if m:
        quarter, year = (m.group(1), m.group(2)) if text.startswith("Q") else (m.group(2), m.group(1))
        return "window", f"{year}-Q{quarter}"
    return None


class Ctx:
    """Everything the sheets share: the stories, the option lists and the ranges in Stories."""

    def __init__(self, title, subtitle, stories, themes, theme_color, timeline, timeline_label, digest_count, windows):
        self.title, self.subtitle = clean(title or "Geotechnical News", 250), clean(subtitle or "", 250)
        self.theme_color = theme_color
        ordered = sorted(stories, key=lambda s: (str(s.get("_digest") or ""), _int(s.get("relevance")),
                                                 _int(s.get("rank"))), reverse=True)
        self.rows = [_story(s, i) for i, s in enumerate(ordered, 1)]
        self.n = len(self.rows)
        self.last = STORY_FIRST + max(self.n, 1) - 1
        known = [str(t) for t in (themes or [])]
        extra = Counter(r["theme"] for r in self.rows if r["theme"] not in known)
        self.themes = list(dict.fromkeys(known + [t for t, _ in extra.most_common()]))
        self.present = [t for t in self.themes if any(r["theme"] == t for r in self.rows)]
        uses = Counter(r["use"] for r in self.rows)
        self.uses = (STANDARD_USES + [u for u, _ in uses.most_common() if u not in STANDARD_USES + [NOT_RATED]]
                     + ([NOT_RATED] if NOT_RATED in uses else []))
        rels = {r["rel"] for r in self.rows}
        self.levels = [5, 4, 3] + [lvl for lvl in (2, 1) if lvl in rels] + ([0] if 0 in rels or not self.rows else [])
        if windows:
            wins = [str(w.get("key") if isinstance(w, dict) else w) for w in windows]
        else:
            wins = sorted({r["window"] for r in self.rows if r["window"]}, reverse=True)
        self.windows = list(dict.fromkeys(clean(w, 60) for w in wins if w))
        self.show_win = bool(windows) or len(self.windows) > 1 or "window" in str(timeline_label).lower()
        self.timeline = [(str(label), _int(count), _timeline_key(label)) for label, count in (timeline or [])]
        self.timeline_label = str(timeline_label or "daily digest")
        dates = {r["digest"] for r in self.rows if r["digest"]}
        if "digest" in self.timeline_label.lower():     # the timeline lists every digest, even empty ones
            dates |= {k[1] for _, _, k in self.timeline if k and k[0] == "date"}
        self.digest_dates = sorted(dates, reverse=True)
        self.digest_list = f"Lists!$K$2:$K${1 + max(len(self.digest_dates), 1)}"   # per-digest counts (Lists)
        self.digest_count = (_int(digest_count) if digest_count is not None
                             else len(dates) if dates else len({r["digest_iso"] for r in self.rows if r["digest_iso"]}))
        by_source = {}
        for r in self.rows:
            by_source.setdefault(r["source"], []).append(r)
        self.sources = sorted(by_source.items(), key=lambda kv: (-len(kv[1]), kv[0].lower()))
        self.source_last = STORY_FIRST + max(len(self.sources), 1) - 1

    def color(self, theme):
        try:
            return hexc(self.theme_color(theme))
        except Exception:     # a theme_color that cannot colour an unknown theme
            return "64748B"

    def rng(self, key):
        return f"Stories!${SC[key]}${STORY_FIRST}:${SC[key]}${self.last}"

    @property
    def flag(self):
        return self.rng("insel")

    def sp(self, *conditions):
        """SUMPRODUCT over the stories in the current selection (exact text matches, no wildcards)."""
        return "SUMPRODUCT(" + self.flag + "".join(f",--({c})" for c in conditions) + ")"

    def cf(self, *pairs):
        """COUNTIFS over the stories in the current selection; pairs are (column key, criterion)."""
        return "COUNTIFS(" + self.flag + ",1" + "".join(f",{self.rng(k)},{v}" for k, v in pairs) + ")"


# ------------------------------------------------------------------ entry point
def write_workbook(path, title, subtitle, stories, themes, theme_color, timeline, timeline_label,
                   digest_count=None, windows=None, papers=None, events=None, deadlines=None, events_asof=None):
    """stories carry _digest/_window/_region (and _country, _storyline) like the site's; timeline is [(label, count)].

    digest_count: number of daily digests in the period (counts digests that found no stories);
    defaults to the digests listed in a daily timeline, else the distinct _digest dates.
    windows: issue-window keys offered by the Issue window selector (default: those in the stories).
    papers: research papers (the digests' "papers", with _digest), listed in a Research sheet.
    events, deadlines: the events calendar's upcoming events and open call-for-papers deadlines as of
    events_asof (events.upcoming() and events.deadlines(); each event may carry _where), in an Events sheet.
    """
    ctx = Ctx(title, subtitle, list(stories or []), themes, theme_color, timeline, timeline_label,
              digest_count, windows)
    ctx.papers = [_paper(p, i) for i, p in enumerate(
        sorted(papers or [], key=lambda p: (_int(p.get("relevance")), str(p.get("published") or "")), reverse=True), 1)]
    ctx.events = [e for e in events or [] if isinstance(e, dict) and e.get("name")]
    ctx.calls = [c for c in deadlines or [] if isinstance(c, dict) and isinstance(c.get("event"), dict)]
    ctx.events_asof = events_asof or dt.date.today()
    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    sheets = {"Start here": start}
    for name in ("Dashboard", "Stories", "By theme", "Research", "Events", "Sources", "Lists"):
        if (name != "Research" or ctx.papers) and (name != "Events" or ctx.events or ctx.calls):
            sheets[name] = wb.create_sheet(name)
    for name, ws in sheets.items():
        ws.sheet_properties.tabColor = TAB_COLORS[name]
        ws.sheet_view.tabSelected = name == "Start here"

    _stories_sheet(sheets["Stories"], ctx)
    _lists_sheet(sheets["Lists"], ctx)
    _sources_sheet(sheets["Sources"], ctx)
    if ctx.papers:
        _research_sheet(sheets["Research"], ctx)
    if "Events" in sheets:
        _events_sheet(sheets["Events"], ctx)
    anchors = _by_theme_sheet(sheets["By theme"], ctx)
    _dashboard(sheets["Dashboard"], ctx)
    _start_sheet(start, ctx, anchors)

    wb.active = 0
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)


# ------------------------------------------------------------------ Stories
def _stories_sheet(ws, ctx):
    put(ws, 1, 1, f"Stories · {ctx.n} in this workbook", f=font(14, True, NAVY), al=Alignment(vertical="center"))
    nav(ws, 1, SCI["theme"], [("← Dashboard", "Dashboard", "A1"), ("↑ Start here", "Start here", "A1"),
                              ("By theme →", "By theme", "A1")])
    put(ws, 2, SCI["theme"], "Gold columns are yours: choose a Status from its dropdown, and fill in Assigned to and "
        "Editor notes. Use ▼ in the header to filter or sort. ✓ in ‘In Dashboard selection’ "
        "(far right) marks the stories the Dashboard is showing.", f=font(9, color=MUTED, italic=True))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 16

    for i, (key, label, width) in enumerate(STORY_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        planning = key in PLANNING
        put(ws, STORY_HEADER, i, label, f=font(10, True, NAVY if planning else WHITE),
            fl=ACCENT if planning else NAVY, al=Alignment(horizontal="center" if width < 13 else "left",
                                                            vertical="center", wrap_text=True))
    ws.row_dimensions[STORY_HEADER].height = 32
    ws.column_dimensions[SC["key"]].hidden = True

    top = Alignment(vertical="top", wrap_text=False)
    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(vertical="top", horizontal="center")
    widths = {key: width for key, _, width in STORY_COLS}
    # a story is in the selection when each selector is "All …" (or cleared) or equals the story's value
    sel = [f'OR({_is_all(SEL_THEME, ALL_THEMES)},${SC["theme"]}{{r}}={DASH}{SEL_THEME})',
           f'OR({_is_all(SEL_USE, ALL_USES)},${SC["use"]}{{r}}={DASH}{SEL_USE})',
           f'N(${SC["rel"]}{{r}})>=Lists!$H$2']
    if ctx.show_win:
        sel.append(f'OR({_is_all(SEL_WIN, ALL_WINDOWS)},${SC["window"]}{{r}}={DASH}{SEL_WIN})')
    flag = f'=IF(${SC["headline"]}{{r}}="",0,IF(AND({",".join(sel)}),1,0))'
    # unique ranking key for the Top 10: relevance, then keyword score, then newest digest, then row
    rank_key = (f'=IF(${SC["insel"]}{{r}}=1,N(${SC["rel"]}{{r}})*1E14+MIN(MAX(N(${SC["score"]}{{r}}),0),9999)*1E10'
           f'+MIN(N(${SC["digest"]}{{r}}),99999)*1E5+99999-MIN(ROW(),99999),"")')

    for idx, s in enumerate(ctx.rows):
        r = STORY_FIRST + idx
        color = ctx.color(s["theme"])
        values = dict(s, status=DEFAULT_STATUS, assigned=None, notes=None)
        for key_, _, _ in STORY_COLS:
            c = SCI[key_]
            if key_ == "insel":
                put(ws, r, c, F(flag.format(r=r)), f=font(10, True, "1E7B4F"), al=center, fmt='"✓";"";""')
            elif key_ == "key":
                put(ws, r, c, F(rank_key.format(r=r)), f=font(9, color=MUTED), fmt="0")
            elif key_ == "headline":
                cell = put(ws, r, c, values["headline"], f=font(10, True, LINK, underline="single" if s["url"] else None), al=wrap)
                if s["url"] and idx < MAX_LINKS:
                    cell.hyperlink = s["url"]
            elif key_ == "theme":
                put(ws, r, c, values["theme"], f=font(9, True, color), fl=tint(color, .85), al=wrap)
            elif key_ == "rel":
                put(ws, r, c, values["rel"], f=font(10, True), al=center, fmt='0;-0;"–"')
            elif key_ == "use":
                put(ws, r, c, values["use"], f=font(9, True, USE_COLORS.get(values["use"], MUTED)), al=wrap)
            elif key_ in ("published", "digest"):
                put(ws, r, c, values[key_], al=top, fmt="dd mmm yyyy")
            elif key_ == "url":
                put(ws, r, c, values["url"], f=font(8, color=MUTED), al=top)
            elif key_ in PLANNING:
                put(ws, r, c, values[key_], f=font(10, key_ == "status"), fl=EDIT_FILL,
                    al=Alignment(vertical="top", wrap_text=key_ == "notes", horizontal="center" if key_ == "status" else None))
            elif key_ in ("no", "score"):
                put(ws, r, c, values[key_], f=font(9, color=MUTED), al=center)
            else:
                v = values[key_]
                put(ws, r, c, v if v != "" else None, f=font(9 if key_ in ("keywords", "edition") else 10,
                                                             italic=key_ == "summary",
                                                             color=MUTED if key_ in ("keywords", "summary") else INK),
                    al=wrap if key_ in WRAPPED else top)
        lines = max(lines_needed(values[k], widths[k]) for k in ("headline", "angle", "original", "summary", "theme"))
        ws.row_dimensions[r].height = min(max(lines, 2), 5) * 12.5 + 3
    if not ctx.rows:
        put(ws, STORY_FIRST, SCI["headline"], None)
        put(ws, STORY_FIRST, SCI["insel"], F(flag.format(r=STORY_FIRST)), fmt='"✓";"";""')
        put(ws, STORY_FIRST, SCI["key"], F(rank_key.format(r=STORY_FIRST)))
        put(ws, STORY_FIRST, SCI["theme"], "No stories in this period yet.", f=font(10, italic=True, color=MUTED))

    last_col = get_column_letter(len(STORY_COLS))
    table = Table(displayName="StoriesTable", ref=f"A{STORY_HEADER}:{last_col}{ctx.last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)

    dv = DataValidation(type="list", formula1="Lists!$E$2:$E$6", allow_blank=True, showDropDown=False,
                        showErrorMessage=True, errorTitle="Status", error="Choose a status from the list.",
                        showInputMessage=True, promptTitle="Status",
                        prompt="Pick: To review, Shortlisted, Commissioned, Published or Not used.")
    dv.add(f"{SC['status']}{STORY_FIRST}:{SC['status']}{ctx.last}")
    ws.add_data_validation(dv)
    _status_colours(ws, f"{SC['status']}{STORY_FIRST}:{SC['status']}{ctx.last}")
    ws.conditional_formatting.add(f"A{STORY_FIRST}:{SC['url']}{ctx.last}",
                                  FormulaRule(formula=[f'${SC["status"]}{STORY_FIRST}="Not used"'],
                                              font=Font(color="9AA3AE")))
    _rel_colours(ws, f"{SC['rel']}{STORY_FIRST}:{SC['rel']}{ctx.last}")

    ws.freeze_panes = f"C{STORY_FIRST}"
    page_setup(ws, f"{STORY_HEADER}:{STORY_HEADER}")
    ws.print_area = f"A1:{SC[PRINT_TO]}{ctx.last}"     # the columns an editor needs on paper


def _is_all(ref, all_text):
    """Formula test: the Dashboard selector at ref is set to its 'All …' value, or has been cleared."""
    return f'OR({DASH}{ref}="{all_text}",{DASH}{ref}="")'


def _status_colours(ws, ref):
    for name, fg, bg, _, _ in STATUSES:
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=[f'"{name}"'], fill=fill(bg),
                                                      font=Font(bold=True, color=fg)))


def _rel_colours(ws, ref):
    for lvl in (5, 4, 3, 2, 1):
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=[str(lvl)], fill=fill(REL_COLORS[lvl]),
                                                      font=Font(bold=True, color=WHITE if lvl >= 3 else INK)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=["0"], font=Font(color="9AA3AE")))


# ------------------------------------------------------------------ Lists (hidden)
def _lists_sheet(ws, ctx):
    ws.sheet_state = "hidden"
    columns = [("Theme selector", [ALL_THEMES] + ctx.themes), ("Use selector", [ALL_USES] + ctx.uses),
               ("Relevance selector", [ANY_REL, 3, 4, 5]), ("Window selector", [ALL_WINDOWS] + ctx.windows),
               ("Status", [s[0] for s in STATUSES])]
    for c, (label, values) in enumerate(columns, 1):
        put(ws, 1, c, label, f=font(10, True))
        for r, v in enumerate(values, 2):
            put(ws, r, c, v)
        ws.column_dimensions[get_column_letter(c)].width = 30
    sel_all = [_is_all(SEL_THEME, ALL_THEMES), _is_all(SEL_USE, ALL_USES), "$H$2=0"]
    if ctx.show_win:
        sel_all.append(_is_all(SEL_WIN, ALL_WINDOWS))
    settings = [("Minimum relevance (number)",
                 F(f"=IF(ISNUMBER({DASH}{SEL_REL}),{DASH}{SEL_REL},IFERROR(VALUE({DASH}{SEL_REL}),0))")),
                ("Daily digests in this workbook", ctx.digest_count),
                ("Stories in this workbook", ctx.n),
                ("No selector set", F(f"=AND({','.join(sel_all)})")),
                ("Stories in the selection", F(f"=COUNTIF({ctx.flag},1)"))]
    put(ws, 1, 7, "Setting", f=font(10, True))
    put(ws, 1, 8, "Value", f=font(10, True))
    for r, (label, value) in enumerate(settings, 2):
        put(ws, r, 7, label)
        put(ws, r, 8, value)
    ws.column_dimensions["G"].width = 30

    put(ws, 1, 10, "Digest date", f=font(10, True))
    put(ws, 1, 11, "Stories in the selection", f=font(10, True))
    for r, d in enumerate(ctx.digest_dates or [None], 2):
        put(ws, r, 10, d, fmt="yyyy-mm-dd")
        put(ws, r, 11, F(f"={ctx.cf(('digest', f'J{r}'))}") if d else 0)
    ws.column_dimensions["J"].width = 14

    # Top 10 helpers: the k-th largest key in Stories, the row it sits on, its link and headline
    for c, label in enumerate(["Rank", "Key", "Row in Stories", "Article URL", "Headline shown"], 13):
        put(ws, 1, c, label, f=font(10, True))
    keys = ctx.rng("key")
    for k in range(1, 11):
        r = k + 1
        put(ws, r, 13, k)
        put(ws, r, 14, F(f'=IFERROR(LARGE({keys},M{r}),"")'), fmt="0")
        put(ws, r, 15, F(f'=IF(N{r}="","",MATCH(N{r},{keys},0))'))
        put(ws, r, 16, F(f'=IF(O{r}="","",INDEX({ctx.rng("url")},O{r})&"")'))
        put(ws, r, 17, F(f'=IF(O{r}="","",IF(LEN(INDEX({ctx.rng("headline")},O{r}))>120,'
                         f'LEFT(INDEX({ctx.rng("headline")},O{r}),118)&"…",INDEX({ctx.rng("headline")},O{r})&""))'))
    ws.column_dimensions["P"].width = ws.column_dimensions["Q"].width = 40
    page_setup(ws, "1:1")


# ------------------------------------------------------------------ Sources
def _sources_sheet(ws, ctx):
    ws.sheet_view.showGridLines = False
    cols = [("Source", 34), ("Stories", 10), ("In current selection", 13), ("Average relevance", 12),
            ("Rated 4 or 5", 10), ("Most frequent theme", 34), ("Themes covered", 10),
            ("Google News editions", 28), ("Latest digest", 13)]
    put(ws, 1, 1, f"Sources · {len(ctx.sources)} outlets", f=font(14, True, NAVY), al=Alignment(vertical="center"))
    nav(ws, 1, 3, [("← Dashboard", "Dashboard", "A1"), ("↑ Start here", "Start here", "A1")])
    put(ws, 2, 1, "One row per outlet, most stories first. ‘In current selection’ follows the Dashboard "
        "selectors; average relevance counts AI-rated stories only.", f=font(9, color=MUTED, italic=True))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 16
    for i, (label, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        put(ws, STORY_HEADER, i, label, f=font(10, True, WHITE), fl=NAVY,
            al=Alignment(horizontal="left" if i in (1, 6, 8) else "center", vertical="center", wrap_text=True))
    ws.row_dimensions[STORY_HEADER].height = 30
    center = Alignment(horizontal="center", vertical="top")
    for i, (src, items) in enumerate(ctx.sources):
        r = STORY_FIRST + i
        rated = [s["rel"] for s in items if s["rel"]]
        top = Counter(s["theme"] for s in items).most_common(1)[0][0]
        editions = Counter(s["edition"] for s in items if s["edition"])
        latest = max((s["digest"] for s in items if s["digest"]), default=None)
        put(ws, r, 1, src, f=font(10, True), al=Alignment(vertical="top", wrap_text=True))
        put(ws, r, 2, len(items), al=center)
        is_source = ctx.rng("source") + f"=$A{r}"
        put(ws, r, 3, F("=" + ctx.sp(is_source)), f=font(10, True, NAVY), al=center)
        put(ws, r, 4, round(sum(rated) / len(rated), 1) if rated else "–", al=center, fmt="0.0")
        put(ws, r, 5, sum(1 for v in rated if v >= 4), al=center)
        put(ws, r, 6, top, f=font(10, True, ctx.color(top)), al=Alignment(vertical="top", wrap_text=True))
        put(ws, r, 7, len({s["theme"] for s in items}), al=center)
        put(ws, r, 8, ", ".join(e for e, _ in editions.most_common()), f=font(9, color=MUTED),
            al=Alignment(vertical="top", wrap_text=True))
        put(ws, r, 9, latest, al=center, fmt="dd mmm yyyy")
        lines = max(lines_needed(src, 34), lines_needed(top, 34), lines_needed(ws.cell(row=r, column=8).value, 28))
        ws.row_dimensions[r].height = min(lines, 4) * 12.5 + 3
    if not ctx.sources:
        put(ws, STORY_FIRST, 1, "No sources yet.", f=font(10, italic=True, color=MUTED))
        put(ws, STORY_FIRST, 3, 0)
    table = Table(displayName="SourcesTable", ref=f"A{STORY_HEADER}:I{ctx.source_last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.conditional_formatting.add(f"B{STORY_FIRST}:B{ctx.source_last}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max", color="7FB3E6",
                                              showValue=True))
    ws.freeze_panes = f"B{STORY_FIRST}"
    page_setup(ws, f"{STORY_HEADER}:{STORY_HEADER}")


# ------------------------------------------------------------------ Research
RESEARCH_COLS = [("No.", 6), ("Paper", 50), ("Theme", 21), ("Use in practice", 10), ("Status", 14),
                 ("Editor notes", 28), ("Why it matters in practice", 44), ("Journal", 26), ("Authors", 22),
                 ("Published", 12), ("Digest date", 12), ("Abstract", 60), ("Keywords", 26), ("DOI link", 30)]


def _research_sheet(ws, ctx):
    """New papers for the Academia-to-Industry pages, best rated for practice first."""
    ws.sheet_view.showGridLines = False
    n = len(ctx.papers)
    last = STORY_FIRST + n - 1
    put(ws, 1, 1, f"Research · {n} new {'paper' if n == 1 else 'papers'}", f=font(14, True, NAVY),
        al=Alignment(vertical="center"))
    nav(ws, 1, 3, [("← Dashboard", "Dashboard", "A1"), ("↑ Start here", "Start here", "A1")])
    put(ws, 2, 3, "Papers from geotechnical journals (via Crossref), best rated for use in practice first. Click a "
        "title to open the paper. Status and Editor notes are yours, as in Stories.", f=font(9, color=MUTED, italic=True))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 16
    for i, (label, width) in enumerate(RESEARCH_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        planning = label in ("Status", "Editor notes")
        put(ws, STORY_HEADER, i, label, f=font(10, True, NAVY if planning else WHITE), fl=ACCENT if planning else NAVY,
            al=Alignment(horizontal="center" if width < 13 else "left", vertical="center", wrap_text=True))
    ws.row_dimensions[STORY_HEADER].height = 32
    top, wrap = Alignment(vertical="top"), Alignment(vertical="top", wrap_text=True)
    center = Alignment(vertical="top", horizontal="center")
    for idx, p in enumerate(ctx.papers):
        r = STORY_FIRST + idx
        color = ctx.color(p["theme"])
        put(ws, r, 1, p["no"], f=font(9, color=MUTED), al=center)
        cell = put(ws, r, 2, p["title"], f=font(10, True, LINK, underline="single" if p["url"] else None), al=wrap)
        if p["url"] and idx < MAX_LINKS:
            cell.hyperlink = p["url"]
        put(ws, r, 3, p["theme"], f=font(9, True, color), fl=tint(color, .85), al=wrap)
        put(ws, r, 4, p["rel"], f=font(10, True), al=center, fmt='0;-0;"–"')
        put(ws, r, 5, DEFAULT_STATUS, f=font(10, True), fl=EDIT_FILL, al=center)
        put(ws, r, 6, None, fl=EDIT_FILL, al=wrap)
        put(ws, r, 7, p["angle"] or None, al=wrap)
        put(ws, r, 8, p["journal"], f=font(10, True), al=wrap)
        put(ws, r, 9, p["authors"] or None, f=font(9), al=wrap)
        put(ws, r, 10, p["published"], al=top, fmt="dd mmm yyyy")
        put(ws, r, 11, p["digest"], al=top, fmt="dd mmm yyyy")
        put(ws, r, 12, p["abstract"] or None, f=font(9, italic=True, color=MUTED), al=wrap)
        put(ws, r, 13, p["keywords"] or None, f=font(9, color=MUTED), al=wrap)
        put(ws, r, 14, p["url"] or None, f=font(8, color=MUTED), al=top)
        lines = max(lines_needed(p["title"], 50), lines_needed(p["angle"], 44), lines_needed(p["journal"], 26))
        ws.row_dimensions[r].height = min(max(lines, 3), 6) * 12.5 + 3
    table = Table(displayName="ResearchTable", ref=f"A{STORY_HEADER}:{get_column_letter(len(RESEARCH_COLS))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    dv = DataValidation(type="list", formula1="Lists!$E$2:$E$6", allow_blank=True, showDropDown=False,
                        showErrorMessage=True, errorTitle="Status", error="Choose a status from the list.")
    dv.add(f"E{STORY_FIRST}:E{last}")
    ws.add_data_validation(dv)
    _status_colours(ws, f"E{STORY_FIRST}:E{last}")
    _rel_colours(ws, f"D{STORY_FIRST}:D{last}")
    ws.freeze_panes = f"C{STORY_FIRST}"
    page_setup(ws, f"{STORY_HEADER}:{STORY_HEADER}")
    ws.print_area = f"A1:K{last}"


# ------------------------------------------------------------------ Events
def _ev_date(iso):
    """A date cell for 'YYYY-MM-DD', the month as text ('Jun 2027') when only the month is known, else None."""
    text = str(iso or "")
    day = as_date(text) if len(text) == 10 else None
    if day:
        return day
    m = re.fullmatch(r"(\d{4})-(\d{2})", text)
    return f"{dt.date(2000, int(m.group(2)), 1):%b} {m.group(1)}" if m and 1 <= int(m.group(2)) <= 12 else None


def _ev_passed(iso, asof):
    """True when a deadline ('YYYY-MM-DD', or 'YYYY-MM' up to the month's end) lies before asof."""
    text = str(iso or "")
    day = as_date(text) if len(text) == 10 else None
    m = re.fullmatch(r"(\d{4})-(\d{2})", text) if not day else None
    if m and 1 <= int(m.group(2)) <= 12:  # the month's last day
        y, mo = int(m.group(1)), int(m.group(2))
        day = dt.date(y + mo // 12, mo % 12 + 1, 1) - dt.timedelta(days=1)
    return bool(day and day < asof)


def _ev_where(e):
    if e.get("_where"):
        return str(e["_where"])
    where = ", ".join(str(x) for x in (e.get("city"), e.get("country")) if x)
    return (f"{where} and online" if where else "Online") if e.get("online") else where


def _ev_name(e):
    short = str(e.get("short") or "")
    name = str(e.get("name") or "")
    return name + (f" ({short})" if short and short.lower() not in name.lower() else "")


def _events_sheet(ws, ctx):
    """The events calendar as it stood when the workbook was made: open deadlines (calls for papers,
    registration), then the upcoming events, each with its website."""
    ws.sheet_view.showGridLines = False
    cols = [("Date", 13), ("Days left", 10), ("Event", 52), ("Kind", 19), ("Place", 26), ("Dates", 22),
            ("Organiser", 28), ("Deadlines", 34), ("Website", 34)]
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    put(ws, 1, 1, f"Events · {len(ctx.events)} upcoming, {len(ctx.calls)} open "
        f"{'deadline' if len(ctx.calls) == 1 else 'deadlines'}", f=font(14, True, NAVY), al=Alignment(vertical="center"))
    nav(ws, 1, 4, [("← Dashboard", "Dashboard", "A1"), ("↑ Start here", "Start here", "A1")])
    put(ws, 2, 1, f"From the website's events calendar on {ctx.events_asof:%d %b %Y}: ISSMGE conferences, other "
        "conferences, courses and webinars, kept by the editors (data/events.yaml) and found by the weekly refresh. "
        "Check the dates on each event's website before you publish them.", f=font(9, color=MUTED, italic=True))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 16
    top, wrap = Alignment(vertical="top"), Alignment(vertical="top", wrap_text=True)
    center = Alignment(vertical="top", horizontal="center")

    def dates(e):
        """'12–14 Apr 2027', '28 Jun – 1 Jul 2027', '1 Nov 2026' or 'Oct 2027', as on the Events page."""
        a, b = _ev_date(e.get("start")), _ev_date(e.get("end"))
        text = lambda d: f"{d.day} {d:%b %Y}" if isinstance(d, dt.date) else str(d or "")
        if not b or a == b or not (isinstance(a, dt.date) and isinstance(b, dt.date)):
            return text(a) if not b or a == b else f"{text(a)} – {text(b)}"
        if (a.year, a.month) == (b.year, b.month):
            return f"{a.day}–{b.day} {b:%b %Y}"
        return f"{a.day} {a:%b} – {text(b)}" if a.year == b.year else f"{text(a)} – {text(b)}"

    def event_cells(r, e):
        # the name is plain wrapped text and the link sits in Website (LibreOffice does not wrap a linked cell)
        kind, url = str(e.get("kind") or ""), safe_url(e.get("url"))
        put(ws, r, 3, clean(_ev_name(e), 400), f=font(10, True), al=wrap)
        put(ws, r, 4, kind or None, f=font(9, True, KIND_COLORS.get(kind, MUTED)), al=wrap)
        put(ws, r, 5, clean(_ev_where(e), 200) or None, al=wrap)
        put(ws, r, 6, dates(e) or None, al=wrap)
        put(ws, r, 7, clean(str(e.get("organiser") or ""), 200) or None, f=font(9), al=wrap)
        cell = put(ws, r, 9, url or None, f=font(9, color=LINK, underline="single" if url else None), al=top)
        if url:
            cell.hyperlink = url
        return max(lines_needed(_ev_name(e), 52), lines_needed(_ev_where(e), 26), lines_needed(e.get("organiser"), 28))

    row = 4
    section(ws, row, 1, len(cols), f"Open deadlines: calls for papers, registration ({len(ctx.calls)})")
    head(ws, row + 1, [(1, 1, "Deadline"), (2, 2, "Days left"), (3, 3, "Event"), (4, 4, "Kind"), (5, 5, "Place"),
                       (6, 6, "Event dates"), (7, 7, "Organiser"), (8, 8, "What is due"), (9, 9, "Website")])
    row += 2
    for c in ctx.calls:
        day = _ev_date(c.get("date"))
        put(ws, row, 1, day, f=font(10, True), al=top, fmt="dd mmm yyyy")
        if isinstance(day, dt.date):  # counted from the day the file is opened
            put(ws, row, 2, F(f"=A{row}-TODAY()"), f=font(10, True, "B4321B"), al=center, fmt='0;"passed";"today"')
        else:
            put(ws, row, 2, "exact date TBC", f=font(9, color=MUTED), al=center)
        lines = event_cells(row, c["event"])
        put(ws, row, 8, clean(str(c.get("label") or "Deadline"), 200), f=font(10, True), al=wrap)
        ws.row_dimensions[row].height = min(max(lines, 1), 4) * 12.5 + 5
        row += 1
    if not ctx.calls:
        put(ws, row, 3, "No deadline is open at the moment.", f=font(10, italic=True, color=MUTED))
        row += 1

    row += 1
    section(ws, row, 1, len(cols), f"Upcoming events ({len(ctx.events)})")
    head(ws, row + 1, [(1, 1, "Starts"), (2, 2, "ISSMGE"), (3, 3, "Event"), (4, 4, "Kind"), (5, 5, "Place"),
                       (6, 6, "Dates"), (7, 7, "Organiser"), (8, 8, "Deadlines"), (9, 9, "Website")])
    row += 2
    for e in ctx.events:
        put(ws, row, 1, _ev_date(e.get("start")), f=font(10, True), al=top, fmt="dd mmm yyyy")
        put(ws, row, 2, "ISSMGE" if e.get("issmge") else None, f=font(9, True, ACCENT), al=center)
        lines = event_cells(row, e)
        # a deadline that has passed is marked, as the Events page strikes it through
        dls = "\n".join(f"{d.get('label', 'Deadline')}: {dates({'start': d.get('date')})}"
                        + (" (passed)" if _ev_passed(d.get("date"), ctx.events_asof) else "")
                        for d in e.get("deadlines") or [] if isinstance(d, dict))
        put(ws, row, 8, clean(dls, 600) or None, f=font(9), al=wrap)
        ws.row_dimensions[row].height = min(max(lines, len(dls.splitlines()), 1), 6) * 12.5 + 5
        row += 1
    if not ctx.events:
        put(ws, row, 3, "No upcoming events in the calendar.", f=font(10, italic=True, color=MUTED))
    ws.freeze_panes = "A4"
    page_setup(ws, "1:2")


# ------------------------------------------------------------------ By theme
def _by_theme_sheet(ws, ctx):
    """One collapsible section per theme; returns {theme: header row} for links."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True, applyStyles=False)
    # the headline is plain wrapped text (LibreOffice does not wrap hyperlinked cells); "Open article" links
    cols = [("No.", 6), ("Relevance", 10), ("Headline", 56), ("Link", 12), ("Suggested use", 17), ("Status", 14),
            ("Source", 20), ("Digest", 12), ("Magazine angle / summary", 50)]
    last_col = len(cols)
    centred = {1, 2, 4, 6, 8}
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    put(ws, 1, 1, "Stories by theme", f=font(14, True, NAVY), al=Alignment(vertical="center"))
    nav(ws, 1, 4, [("\u2190 Dashboard", "Dashboard", "A1"), ("\u2191 Start here", "Start here", "A1"),
                   ("Stories \u2192", "Stories", "A1")])
    put(ws, 2, 1, "Read-only view, best stories first. Click \u2212 / + in the margin on the left to collapse or "
        "expand a theme (or the 1 / 2 buttons for all). Status comes from the Stories sheet.",
        f=font(9, color=MUTED, italic=True))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 26, 16
    status_rng, no_rng = ctx.rng("status"), ctx.rng("no")
    anchors, row = {}, 4
    if not ctx.present:
        put(ws, row, 1, "No stories in this period yet.", f=font(10, italic=True, color=MUTED))
    for theme in ctx.present:
        items = sorted((s for s in ctx.rows if s["theme"] == theme),
                       key=lambda s: (-s["rel"], -s["score"], -(s["digest"] or dt.date.min).toordinal(), s["no"]))
        color = ctx.color(theme)
        anchors[theme] = row
        n5 = sum(1 for s in items if s["rel"] == 5)
        leads = sum(1 for s in items if s["use"] == "Feature lead")
        text = (f"{theme}   \u00b7   {len(items)} {'story' if len(items) == 1 else 'stories'}   \u00b7   "
                f"{n5} rated 5 / 5   \u00b7   {leads} feature {'lead' if leads == 1 else 'leads'}")
        span(ws, row, 1, last_col, text, f=font(12, True, WHITE), fl=color, al=Alignment(vertical="center", indent=1))
        ws.row_dimensions[row].height = 26
        row += 1
        for c, (label, _) in enumerate(cols, 1):
            put(ws, row, c, label, f=font(9, True, color), fl=tint(color, .86),
                al=Alignment(horizontal="center" if c in centred else "left", vertical="center"),
                border=Border(bottom=Side(style="thin", color=color)))
        ws.row_dimensions[row].outline_level = 1
        row += 1
        first = row
        for s in items:
            top = Alignment(vertical="top", wrap_text=True)
            mid = Alignment(horizontal="center", vertical="top")
            put(ws, row, 1, s["no"], f=font(9, color=MUTED), al=mid)
            put(ws, row, 2, s["rel"], f=font(10, True), al=mid, fmt='0;-0;"\u2013"')
            put(ws, row, 3, s["headline"], f=font(10, True, INK), al=top)
            if s["url"]:
                put(ws, row, 4, "Open article", f=font(9, color=LINK, underline="single"), al=mid).hyperlink = s["url"]
            put(ws, row, 5, s["use"], f=font(9, True, USE_COLORS.get(s["use"], MUTED)), al=top)
            put(ws, row, 6, F(f'=IFERROR(INDEX({status_rng},MATCH($A{row},{no_rng},0))&"","")'), al=mid)
            put(ws, row, 7, s["source"], f=font(9), al=top)
            put(ws, row, 8, s["digest"], f=font(9), al=mid, fmt="dd mmm yyyy")
            put(ws, row, 9, s["angle"] or s["summary"], f=font(9, color=INK if s["angle"] else MUTED,
                                                                   italic=not s["angle"]), al=top)
            for c in range(1, last_col + 1):
                ws.cell(row=row, column=c).border = Border(bottom=THIN)
            lines = max(lines_needed(s["headline"], 56), lines_needed(s["angle"] or s["summary"], 50),
                        lines_needed(s["source"], 20))
            ws.row_dimensions[row].height = min(max(lines, 1), 5) * 12.5 + 4
            ws.row_dimensions[row].outline_level = 1
            row += 1
        _rel_colours(ws, f"B{first}:B{row - 1}")
        _status_colours(ws, f"F{first}:F{row - 1}")
        ws.row_dimensions[row].outline_level = 1
        ws.row_dimensions[row].height = 10
        row += 1
    ws.sheet_format.outlineLevelRow = 1
    ws.freeze_panes = "A3"
    page_setup(ws, "1:2")
    return anchors


# ------------------------------------------------------------------ Dashboard
def _dashboard(ws, ctx):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for c in range(2, 20):
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.column_dimensions["T"].width = 2
    B, S = 2, 19

    # ---- banner
    for c in range(1, 21):
        ws.cell(row=1, column=c).fill = fill(NAVY)
        ws.cell(row=2, column=c).fill = fill(NAVY2)
        ws.cell(row=3, column=c).fill = fill(ACCENT)
    span(ws, 1, B, S, ctx.title, f=font(18, True, WHITE), fl=NAVY, al=Alignment(vertical="center"))
    span(ws, 2, B, S, f"{ctx.subtitle} · generated {dt.date.today():%d %b %Y}".strip(" ·"),
         f=font(10, color=WHITE, italic=True), fl=NAVY2, al=Alignment(vertical="center"))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 34, 20, 4

    # ---- selectors
    label_font, input_font = font(8, True, MUTED), font(11, True, NAVY)
    selectors = [(2, 6, "THEME", SEL_THEME, ALL_THEMES, f"Lists!$A$2:$A${2 + len(ctx.themes)}",
                  "Choose one theme, or All themes."),
                 (7, 10, "SUGGESTED USE", SEL_USE, ALL_USES, f"Lists!$B$2:$B${2 + len(ctx.uses)}",
                  "Choose an editorial use, or All uses."),
                 (11, 13, "MINIMUM RELEVANCE (OUT OF 5)", SEL_REL, ANY_REL, "Lists!$C$2:$C$5",
                  "Show stories rated at least 3, 4 or 5 by the AI triage, or Any.")]
    if ctx.show_win:
        selectors.append((14, 16, "ISSUE WINDOW", SEL_WIN, ALL_WINDOWS, f"Lists!$D$2:$D${2 + len(ctx.windows)}",
                          "Choose one quarterly issue window, or All windows."))
    for c1, c2, label, ref, default, source, prompt in selectors:
        span(ws, 4, c1, c2, label, f=label_font, al=Alignment(vertical="bottom"))
        span(ws, 5, c1, c2, default, f=input_font, fl=INPUT_FILL,
             al=Alignment(horizontal="left", vertical="center", indent=1))
        box(ws, 5, c1, 5, c2, Side(style="medium", color=ACCENT))
        dv = DataValidation(type="list", formula1=source, allow_blank=False, showErrorMessage=True,
                            errorTitle=label.split(" (")[0].capitalize(), error="Please pick a value from the list.",
                            showInputMessage=True, promptTitle=label.split(" (")[0].capitalize(), prompt=prompt)
        dv.add(ref.replace("$", ""))
        ws.add_data_validation(dv)
    if not ctx.show_win:
        span(ws, 4, 14, 16, "ISSUE WINDOW", f=label_font, al=Alignment(vertical="bottom"))
        span(ws, 5, 14, 16, ", ".join(ctx.windows) or "–", f=font(11, True, MUTED), fl="EEF1F4",
             al=Alignment(vertical="center", indent=1))
    span(ws, 4, 17, 19, "Tip: click a yellow box, then its ▼ arrow, to choose. Pick “All …” or "
         "“Any” to clear it. Everything below updates.", f=font(8, color=MUTED, italic=True), fl=PAPER,
         al=Alignment(vertical="center", wrap_text=True, indent=1), r2=5)
    ws.row_dimensions[4].height, ws.row_dimensions[5].height = 16, 26
    span(ws, 6, B, S, F(f'="Showing "&Lists!$H$6&" of "&Lists!$H$4&" stories"&IF(Lists!$H$5," – no selector set.",'
                        f'" for the selection above. Set every box back to All / Any to see everything.")'),
         f=font(9, color=MUTED, italic=True), al=Alignment(vertical="center"))
    ws.row_dimensions[6].height, ws.row_dimensions[7].height = 18, 8

    # ---- key figures
    planned = "+".join(ctx.cf(("status", q(s))) for s in ("Shortlisted", "Commissioned", "Published"))
    tiles = [("STORIES", F(f"=Lists!$H$6"), F('="of "&Lists!$H$4&" in this workbook"')),
             (F('=IF(Lists!$H$5,"DAILY DIGESTS","DIGESTS WITH MATCHES")'),
              F(f'=IF(Lists!$H$5,Lists!$H$3,COUNTIF({ctx.digest_list},">0"))'),
              F('=IF(Lists!$H$5,"scans in this period","of "&Lists!$H$3&" scans in this period")')),
             ("RATED 5 / 5", F(f"={ctx.cf(('rel', 5))}"), F(f'="rated 4 / 5: "&{ctx.cf(("rel", 4))}')),
             ("FEATURE LEADS", F("=" + ctx.cf(("use", q("Feature lead")))),
              F('="news briefs: "&' + ctx.cf(("use", q("News brief"))))),
             ("PLANNED", F("=" + planned),
              F('="shortlisted or later · to review: "&' + ctx.cf(("status", q(DEFAULT_STATUS))))),
             ("SOURCES", F(f'=COUNTIF(Sources!$C${STORY_FIRST}:$C${ctx.source_last},">0")'),
              F(f'="of {len(ctx.sources)} outlets in this workbook"'))]
    for i, (label, value, caption) in enumerate(tiles):
        c1, c2 = B + 3 * i, B + 3 * i + 2
        color = KPI_COLORS[i]
        sep = Border(left=Side(style="thick", color=WHITE), right=Side(style="thick", color=WHITE))
        span(ws, 8, c1, c2, value, f=font(22, True, WHITE), fl=color, fmt="#,##0",
             al=Alignment(horizontal="center", vertical="center"), border=sep)
        span(ws, 9, c1, c2, label, f=font(8, True, color), fl=tint(color, .86),
             al=Alignment(horizontal="center", vertical="center"), border=sep)
        span(ws, 10, c1, c2, caption, f=font(8, color=MUTED), fl=tint(color, .86),
             al=Alignment(horizontal="center", vertical="center"), border=sep)
    ws.row_dimensions[8].height, ws.row_dimensions[9].height, ws.row_dimensions[10].height = 36, 15, 15
    ws.row_dimensions[11].height = 12

    # ---- Top 10
    section(ws, 12, B, S, "Top 10 stories for this selection")
    head(ws, 13, [(2, 2, "#"), (3, 8, "Headline (click to open the article)"), (9, 9, "Digest"),
                  (10, 12, "Theme"), (13, 13, "Relevance"), (14, 15, "Suggested use"), (16, 17, "Source"),
                  (18, 18, "Story No."), (19, 19, "Status")])
    row_of = lambda k: f"Lists!$O${k + 1}"
    for k in range(1, 11):
        r = 13 + k
        def pick(key, k=k, text=True):     # text lookups get &"" so an emptied cell shows blank, not 0
            suffix = '&""' if text else ""
            return F(f'=IF({row_of(k)}="","",INDEX({ctx.rng(key)},{row_of(k)}){suffix})')
        band = PAPER if k % 2 == 0 else WHITE
        bottom = Border(bottom=THIN)
        put(ws, r, 2, F(f'=IF({row_of(k)}="","",{k})'), f=font(12, True, ACCENT), fl=band,
            al=Alignment(horizontal="center", vertical="center"), border=bottom)
        empty = q("No stories match this selection – choose “All …” or “Any” in a box "
                  "above.") if k == 1 else '""'
        # Excel's HYPERLINK() takes at most 255 characters: longer article links (most Google News ones)
        # jump to the story's row in Stories, where the headline itself links to the article
        url, shown = f"Lists!$P${k + 1}", f"Lists!$Q${k + 1}"
        to_row = f"\"#'Stories'!{SC['headline']}\"&({row_of(k)}+{STORY_FIRST - 1})"
        span(ws, r, 3, 8, F(f'=IF({row_of(k)}="",{empty},IF(AND({url}<>"",LEN({url})<=255),'
                            f'HYPERLINK({url},{shown}),HYPERLINK({to_row},{shown})))'),
             f=font(10, True, MUTED), fl=band, border=bottom,
             al=Alignment(vertical="center", wrap_text=True, indent=1))
        put(ws, r, 9, pick("digest", text=False), f=font(9), fl=band, fmt="dd mmm yy;;", border=bottom,
            al=Alignment(horizontal="center", vertical="center"))
        span(ws, r, 10, 12, pick("theme"), f=font(9, True), fl=band, border=bottom,
             al=Alignment(vertical="center", wrap_text=True, indent=1))
        put(ws, r, 13, pick("rel", text=False), f=font(10, True), fl=band, fmt='0;-0;"–"', border=bottom,
            al=Alignment(horizontal="center", vertical="center"))
        span(ws, r, 14, 15, pick("use"), f=font(9), fl=band, border=bottom,
             al=Alignment(vertical="center", wrap_text=True, indent=1))
        span(ws, r, 16, 17, pick("source"), f=font(9), fl=band, border=bottom,
             al=Alignment(vertical="center", wrap_text=True, indent=1))
        put(ws, r, 18, pick("no", text=False), f=font(9, color=MUTED), fl=band, border=bottom,
            al=Alignment(horizontal="center", vertical="center"))
        put(ws, r, 19, pick("status"), f=font(8), fl=band, border=bottom,
            al=Alignment(horizontal="center", vertical="center"))
        ws.row_dimensions[r].height = 30
    for i, theme in enumerate(ctx.themes):    # theme cell coloured by the theme it shows
        color = ctx.color(theme)
        ws.conditional_formatting.add("J14:J23", FormulaRule(formula=[f"$J14=Lists!$A${3 + i}"],
                                                              fill=fill(tint(color, .85)), font=Font(bold=True, color=color)))
    # headlines look like links only when there is a story (the "no stories" message stays plain)
    ws.conditional_formatting.add("C14:C23", FormulaRule(formula=['Lists!$O2<>""'],
                                                          font=Font(bold=True, color=LINK, underline="single")))
    _rel_colours(ws, "M14:M23")
    _status_colours(ws, "S14:S23")
    put(ws, 24, 3, "Ranked by relevance, then keyword score, then newest digest. Find a story in the Stories "
        "sheet by its Story No.", f=font(8, color=MUTED, italic=True), al=Alignment(vertical="center"))
    row = 26
    starts = [row]      # first row of each block below the Top 10, for page breaks

    # ---- stories by theme
    sel_total = "Lists!$H$6"
    section(ws, row, B, 10, "Stories by theme")
    head(ws, row + 1, [(2, 5, "Theme"), (6, 6, "Stories"), (7, 7, "Rated 5"), (8, 8, "Rated 4"),
                       (9, 9, "Feature leads"), (10, 10, "Share")], height=28)
    first = row + 2
    for i, theme in enumerate(ctx.themes):
        r = first + i
        color = ctx.color(theme)
        is_theme = f"{ctx.rng('theme')}=$B{r}"
        span(ws, r, 2, 5, theme, f=font(10, True, color), fl=tint(color, .86),
             al=Alignment(vertical="center", indent=1),
             border=Border(left=Side(style="thick", color=color), bottom=THIN))
        put(ws, r, 6, F(f"={ctx.sp(is_theme)}"), f=font(10, True), al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        put(ws, r, 7, F(f"={ctx.sp(is_theme, ctx.rng('rel') + '=5')}"), al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        put(ws, r, 8, F(f"={ctx.sp(is_theme, ctx.rng('rel') + '=4')}"), al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        put(ws, r, 9, F("=" + ctx.sp(is_theme, ctx.rng("use") + "=" + q("Feature lead"))),
            al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        put(ws, r, 10, F(f"=IF({sel_total}=0,0,F{r}/{sel_total})"), f=font(9, color=MUTED), fmt="0%",
            al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        ws.row_dimensions[r].height = 18
    last = first + len(ctx.themes) - 1
    table_pt = 22 + 28 + 18 * len(ctx.themes)
    height_cm = max(table_pt / 28.35, 6.5)
    bar_chart(ws, f"L{row}", Reference(ws, min_col=2, min_row=first, max_row=last),
              [Reference(ws, min_col=6, min_row=row + 1, max_row=last)], title="Stories by theme (this selection)",
              width=16.8, height=height_cm, horizontal=True, colors=[ctx.color(t) for t in ctx.themes])
    row = _after(ws, last + 1, row, height_cm)

    # ---- three panels: suggested use, relevance, status
    # (name, formula given the label cell, colour); the use table matches on its label cell
    panels = [
        ("Suggested editorial use", "Use",
         [(u, lambda cell: "=" + ctx.sp(ctx.rng("use") + "=" + cell), USE_COLORS.get(u, MUTED)) for u in ctx.uses]),
        ("Geotechnical relevance", "Relevance",
         [(REL_LABELS[lvl], lambda cell, lvl=lvl: "=" + ctx.cf(("rel", lvl)), REL_COLORS[lvl]) for lvl in ctx.levels]),
        ("Planning status", "Status",
         [(name, lambda cell, name=name: "=" + ctx.cf(("status", q(name))), chart_color)
          for name, _, _, chart_color, _ in STATUSES]),
    ]
    longest = max(len(p[2]) for p in panels)
    starts.append(row)
    for p, (heading, label, items) in enumerate(panels):
        c0 = B + 6 * p
        section(ws, row, c0, c0 + 4, heading)
        head(ws, row + 1, [(c0, c0 + 2, label), (c0 + 3, c0 + 3, "Stories"), (c0 + 4, c0 + 4, "Share")])
        for i, (name, formula, color) in enumerate(items):
            r = row + 2 + i
            lc = get_column_letter(c0)
            span(ws, r, c0, c0 + 2, name, f=font(10, True, color if p != 2 else STATUSES[i][1]),
                 al=Alignment(vertical="center", indent=1), border=Border(bottom=THIN),
                 fl=STATUSES[i][2] if p == 2 else None)
            put(ws, r, c0 + 3, F(formula(f"${lc}{r}")), f=font(10, True), al=Alignment(horizontal="center"), border=Border(bottom=THIN))
            cnt = f"{get_column_letter(c0 + 3)}{r}"
            put(ws, r, c0 + 4, F(f"=IF({sel_total}=0,0,{cnt}/{sel_total})"), f=font(9, color=MUTED), fmt="0%",
                al=Alignment(horizontal="center"), border=Border(bottom=THIN))
            ws.row_dimensions[r].height = 18
        chart_row = row + 3 + longest
        bar_chart(ws, f"{get_column_letter(c0)}{chart_row}",
                  Reference(ws, min_col=c0, min_row=row + 2, max_row=row + 1 + len(items)),
                  [Reference(ws, min_col=c0 + 3, min_row=row + 1, max_row=row + 1 + len(items))],
                  title=heading, width=11.2, height=6.2, horizontal=p == 0,
                  colors=[it[2] for it in items])
    row = _after(ws, row + 3 + longest, row + 3 + longest, 6.2)

    # ---- heat map
    uses = ctx.uses
    end_col = 6 + len(uses)
    starts.append(row)
    section(ws, row, B, end_col, "Theme × suggested use (heat map for this selection)")
    head(ws, row + 1, [(2, 5, "Theme")] + [(6 + j, 6 + j, u) for j, u in enumerate(uses)] + [(end_col, end_col, "Total")],
         height=42)
    first = row + 2
    for i, theme in enumerate(ctx.themes):
        r = first + i
        color = ctx.color(theme)
        span(ws, r, 2, 5, theme, f=font(10, True, color), al=Alignment(vertical="center", indent=1),
             border=Border(left=Side(style="thick", color=color), bottom=THIN))
        for j in range(len(uses)):
            col = get_column_letter(6 + j)
            put(ws, r, 6 + j, F(f"={ctx.sp(ctx.rng('theme') + f'=$B{r}', ctx.rng('use') + f'={col}${row + 1}')}"),
                al=Alignment(horizontal="center"), border=Border(bottom=THIN, right=Side(style="thin", color=WHITE)))
        put(ws, r, end_col, F(f"=SUM(F{r}:{get_column_letter(end_col - 1)}{r})"), f=font(10, True),
            al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        ws.row_dimensions[r].height = 18
    last = first + len(ctx.themes) - 1
    ws.conditional_formatting.add(f"F{first}:{get_column_letter(end_col - 1)}{last}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                 mid_type="percentile", mid_value=50, mid_color="F6E3A6",
                                                 end_type="max", end_color="D38B00"))
    row = last + 3

    # ---- timeline
    label = ctx.timeline_label
    starts.append(row)
    section(ws, row, B, S, f"Stories per {label}")
    if not ctx.timeline:
        put(ws, row + 1, 2, f"No {label}s yet.", f=font(10, italic=True, color=MUTED))
        _finish_dashboard(ws, starts, row + 2)
        return
    has_sel = all(k for _, _, k in ctx.timeline)
    head(ws, row + 1, [(2, 5, label[:1].upper() + label[1:]), (6, 6, "All stories")]
         + ([(7, 7, "This selection")] if has_sel else []), height=28)
    first = row + 2
    for i, (text, count, k) in enumerate(ctx.timeline):
        r = first + i
        span(ws, r, 2, 5, text, f=font(10), al=Alignment(vertical="center", indent=1), border=Border(bottom=THIN))
        put(ws, r, 6, count, f=font(10, color=MUTED), al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        if has_sel:
            criterion = f"DATE({k[1].year},{k[1].month},{k[1].day})" if k[0] == "date" else q(k[1])
            put(ws, r, 7, F(f"={ctx.cf(('digest' if k[0] == 'date' else 'window', criterion))}"), f=font(10, True),
                al=Alignment(horizontal="center"), border=Border(bottom=THIN))
        ws.row_dimensions[r].height = 18
    last = first + len(ctx.timeline) - 1
    height_cm = min(max((28 + 18 * len(ctx.timeline)) / 28.35, 6.5), 10)
    bar_chart(ws, f"I{row + 1}", Reference(ws, min_col=2, min_row=first, max_row=last),
              [Reference(ws, min_col=c, min_row=row + 1, max_row=last) for c in ((6, 7) if has_sel else (6,))],
              title=f"Stories per {label}", width=20.6, height=height_cm, legend=has_sel,
              series_colors=["D9DEE3", ACCENT] if has_sel else [ACCENT], overlap=100 if has_sel else None, gap=40)
    _finish_dashboard(ws, starts, _after(ws, last + 1, row + 1, height_cm))


def _after(ws, next_row, top_row, height_cm):
    """First free row below both the table that ends before next_row and a chart of height_cm placed at top_row."""
    needed_pt, used, r = height_cm * 28.35, 0.0, top_row
    while used < needed_pt:
        used += ws.row_dimensions[r].height or 15
        r += 1
    return max(next_row, r) + 1


def _finish_dashboard(ws, starts, end):
    """Freeze the selectors and start a printed page at a block boundary whenever the next block would not fit."""
    ws.freeze_panes = "A7"
    page_setup(ws, "1:6")
    height = lambda a, b: sum(ws.row_dimensions[r].height or 15 for r in range(a, b))
    capacity = 590      # points of rows per A4 landscape page at the fit-to-width scale, below the repeated rows 1-6
    used = height(7, starts[0])
    for start, stop in zip(starts, starts[1:] + [end]):
        block = height(start, stop)
        if used + block > capacity and used > 0:
            ws.row_breaks.append(Break(id=start - 1))
            used = 0
        used += block


# ------------------------------------------------------------------ Start here
def _start_sheet(ws, ctx, anchors):
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 38, "C": 70, "D": 24, "E": 2}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in range(1, 6):
        ws.cell(row=1, column=c).fill = fill(NAVY)
        ws.cell(row=2, column=c).fill = fill(NAVY2)
        ws.cell(row=3, column=c).fill = fill(ACCENT)
    span(ws, 1, 2, 4, ctx.title, f=font(18, True, WHITE), fl=NAVY, al=Alignment(vertical="center"))
    span(ws, 2, 2, 4, f"{ctx.subtitle} · generated {dt.date.today():%d %b %Y}".strip(" ·"),
         f=font(10, color=WHITE, italic=True), fl=NAVY2, al=Alignment(vertical="center"))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 34, 20, 4
    dates = [r["digest"] for r in ctx.rows if r["digest"]]
    period = (f", {min(dates):%d %b %Y} to {max(dates):%d %b %Y}" if dates else "")
    span(ws, 5, 2, 4, f"This workbook holds {ctx.n} {'story' if ctx.n == 1 else 'stories'} from {ctx.digest_count} "
         f"daily {'digest' if ctx.digest_count == 1 else 'digests'}{period}, in {len(ctx.present)} "
         f"{'theme' if len(ctx.present) == 1 else 'themes'} and from {len(ctx.sources)} "
         f"{'source' if len(ctx.sources) == 1 else 'sources'}.", f=font(11, color=INK),
         al=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[5].height = 30
    row = 7

    def block(title, rows):
        nonlocal row
        section(ws, row, 2, 4, title)
        row += 1
        for left, text, style in rows:
            cell = put(ws, row, 2, left, f=style.get("font", font(10, True, NAVY)), fl=style.get("fill"),
                       al=Alignment(vertical="top", horizontal=style.get("h", "left"), indent=style.get("indent", 0)))
            if style.get("link"):
                link_to(cell, *style["link"])
            span(ws, row, 3, 4, text, f=style.get("text_font", font(10)), al=Alignment(vertical="top", wrap_text=True))
            ws.row_dimensions[row].height = max(lines_needed(text, 100), 1) * 13 + 5
            for c in range(2, 5):
                ws.cell(row=row, column=c).border = Border(bottom=THIN)
            row += 1
        row += 1

    link_font = font(10, True, LINK, underline="single")
    block("What is in this workbook", [
        ("Dashboard →", "Key figures, the Top 10 stories, charts by theme, use, relevance and planning status, a "
         "theme × use heat map and a timeline. Choose what it shows with the yellow selectors at the top.",
         {"font": link_font, "link": ("Dashboard", "A1")}),
        ("Stories →", "Every story as a filterable Excel table, newest digest first; click a headline to open the "
         "article. The gold columns – Status, Assigned to and Editor notes – are for you to fill in. Further right: "
         "where each story happened (Country, Place) and the Developing story it belongs to.",
         {"font": link_font, "link": ("Stories", "A1")}),
        ("By theme →", "A read-only reading view: one section per theme, best stories first. Collapse or expand "
         "sections with the − / + buttons in the left margin.", {"font": link_font, "link": ("By theme", "A1")}),
    ] + ([
        ("Research →", f"{len(ctx.papers)} new {'paper' if len(ctx.papers) == 1 else 'papers'} from geotechnical "
         "journals for the Academia-to-Industry pages, best rated for use in practice first, with why each matters "
         "and a link to the paper. Status and Editor notes work as in Stories.",
         {"font": link_font, "link": ("Research", "A1")}),
    ] if ctx.papers else []) + ([
        ("Events →", f"The events calendar on {ctx.events_asof:%d %b %Y}: {len(ctx.calls)} open "
         f"{'deadline' if len(ctx.calls) == 1 else 'deadlines'} (calls for papers, registration) with the days left, "
         f"and {len(ctx.events)} upcoming "
         f"{'event' if len(ctx.events) == 1 else 'events'} (ISSMGE conferences, other conferences, courses and "
         "webinars), each linked to its website.", {"font": link_font, "link": ("Events", "A1")}),
    ] if ctx.events or ctx.calls else []) + [
        ("Sources →", "The outlets that supplied stories: how many, average relevance, main theme and how many "
         "match the current Dashboard selection.", {"font": link_font, "link": ("Sources", "A1")}),
    ])
    block("How to use the Dashboard selectors", [
        ("1", "Open the Dashboard. The yellow boxes at the top are selectors: Theme, Suggested use, Minimum "
         "relevance" + (" and Issue window." if ctx.show_win else " (the issue window is fixed in this workbook)."),
         {"h": "center"}),
        ("2", "Click a yellow box, then the ▼ arrow that appears beside it, and pick a value. Choose “All "
         "…” or “Any” to clear that selector.", {"h": "center"}),
        ("3", "The key figures, tables, charts, heat map and the Top 10 list recalculate at once. Headlines in the "
         "Top 10 open the article.", {"h": "center"}),
        ("4", "In Stories, set each story’s Status as you plan the issue. The Dashboard’s Planning status "
         "table and chart, and the PLANNED figure, count them straight away.", {"h": "center"}),
        ("5", "Filters in the Stories table (▼ in its header) only hide rows; the Dashboard follows the "
         "selectors. To list exactly what the Dashboard shows, filter the ‘In Dashboard selection’ "
         "column of Stories to ✓.", {"h": "center"}),
    ])
    block("Planning status (Stories › Status)", [
        (name, meaning, {"font": font(10, True, fg), "fill": bg, "h": "center"})
        for name, fg, bg, _, meaning in STATUSES])
    block("Relevance scale (AI triage)", [
        (REL_LABELS[lvl] if lvl else "–", REL_MEANING[lvl],
         {"font": font(10, True, WHITE if lvl >= 3 else INK), "fill": REL_COLORS[lvl], "h": "center"})
        for lvl in (5, 4, 3, 2, 1, 0)])
    theme_rows = []
    for theme in ctx.present:
        items = [s for s in ctx.rows if s["theme"] == theme]
        n5 = sum(1 for s in items if s["rel"] == 5)
        leads = sum(1 for s in items if s["use"] == "Feature lead")
        color = ctx.color(theme)
        theme_rows.append((theme, f"{len(items)} {'story' if len(items) == 1 else 'stories'} · {n5} rated 5 / 5 "
                                  f"· {leads} feature {'lead' if leads == 1 else 'leads'}",
                           {"font": font(10, True, color, underline="single"), "fill": tint(color, .86),
                            "link": ("By theme", f"A{anchors[theme]}"), "indent": 1}))
    missing = [t for t in ctx.themes if t not in ctx.present]
    if missing:
        theme_rows.append(("No stories yet", ", ".join(missing), {"font": font(10, True, MUTED)}))
    block("Jump to a theme (opens its section in By theme)", theme_rows or
          [("No stories yet", "Themes appear here once a scan has found stories.", {"font": font(10, True, MUTED)})])
    block("Good to know", [
        ("Blank figures?", "The Dashboard is calculated when the file opens. If it shows blanks or zeros in Excel, "
         "click Enable Editing in the yellow bar at the top (files from the internet open in Protected View); in "
         "LibreOffice press Ctrl+Shift+F9.", {"font": font(10, True, INK)}),
        ("Your edits", "This file is generated again after every scan. Status, Assigned to and Editor notes are kept "
         "only in the copy you save, so keep one working copy per issue.", {"font": font(10, True, INK)}),
        ("Lists sheet", "A hidden sheet holds the values offered by the dropdowns and some helper cells. There is no "
         "need to change it.", {"font": font(10, True, INK)}),
    ])
    page_setup(ws, "1:3")
