"""
Word report on the website: current status, an automatic review of the latest
scan, source health, archive figures, regional coverage on the world map,
developing stories, research papers, issue briefs, the events calendar, the
e-mails sent, the change log and open items. Regenerated after every scan, so it
is always up to date. The change log and open items come from data/changelog.yaml.
"""
import datetime as dt
import json
from collections import Counter

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, ACCENT, MUTED = "16324F", "C9A227", "5B6573"


def rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def shade(cell, hexstr):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexstr.lstrip("#").upper())
    tc_pr.append(shd)


def tint(hexstr, amount=0.85):
    c = hexstr.lstrip("#")
    return "".join(f"{round(int(c[i:i + 2], 16) + (255 - int(c[i:i + 2], 16)) * amount):02X}" for i in (0, 2, 4))


def table(doc, headers, rows, widths, first_col_colors=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(9.5), rgb("FFFFFF")
        shade(cell, NAVY)
    for r, values in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(values):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
            if i == 0 and first_col_colors:
                run.bold = True
                run.font.color.rgb = rgb(first_col_colors[r])
                shade(cells[i], tint(first_col_colors[r]))
            elif r % 2:
                shade(cells[i], "F3F5F7")
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = rgb(NAVY if level == 1 else "23466B")
    return h


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Geotechnical News Desk · Site report · page ")
    run.font.size, run.font.color.rgb = Pt(8), rgb(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def source_advice(error):
    e = error.lower()
    if "403" in e or "405" in e:
        return "The site blocks automated readers. Look for its RSS feed address, or rely on the Google News searches."
    if "no rss" in e:
        return "No feed at this address. Find the site's RSS feed (often /feed/ or an RSS icon) and update config.yaml."
    if "404" in e:
        return "The address has moved. Update it in config.yaml."
    if "timeout" in e or "max retries" in e:
        return "The site did not answer. If it fails again on the next scans, replace or remove it."
    return "Check the address in config.yaml."


def review_findings(latest, cfg, windows):
    """Automatic review of the latest scan: (level, text) pairs."""
    out = []
    if not latest:
        return [("info", "No scan has run yet.")]
    stories = latest["stories"]
    if latest.get("ai_used"):
        out.append(("ok", "Claude triage ran: stories are rated, themed, given an angle and translated."))
    else:
        note = next((n for n in latest.get("notes", []) if "AI triage" in n), "")
        out.append(("warn", "Claude triage did not run, so stories are ranked by keywords only. "
                    + (note or "Add the ANTHROPIC_API_KEY secret to enable it.")))
    feeds = len(cfg.get("feeds", []))
    failed = len(latest.get("failures", []))
    out.append(("warn" if failed > feeds / 4 else "ok",
                f"{feeds - failed} of {feeds} websites were read successfully; {failed} failed (see Source health)."))
    if stories:
        themes = Counter(s["theme"] for s in stories)
        top, n = themes.most_common(1)[0]
        share = n / len(stories)
        out.append(("warn" if share > 0.6 else "ok",
                    f"Theme balance: {len(themes)} themes covered; the largest is {top} with {n} of "
                    f"{len(stories)} stories ({share:.0%})."))
        sources = Counter(s["source"] for s in stories)
        out.append(("info", f"{len(sources)} different outlets contributed; the most frequent was "
                            f"{sources.most_common(1)[0][0]} ({sources.most_common(1)[0][1]} stories)."))
        foreign = sum(1 for s in stories if s.get("region") and s.get("region") not in ("GB", "US", "IN", "AU", "ZA", "SG", "AE"))
        out.append(("info", f"{foreign} stories came from non-English Google News editions."))
        if latest.get("ai_used"):
            strong = sum(1 for s in stories if s.get("relevance", 0) >= 4)
            out.append(("info", f"{strong} stories were rated 4 or 5 out of 5."))
    else:
        out.append(("warn", "The latest scan selected no stories."))
    return out


def feature(module, name, *args, default=None, **kw):
    """module.name(*args) from a feature module (geomap, storylines, events, notify), or default when it is
    missing or fails: its section of the report then says so, and the report is still written."""
    try:
        fn = getattr(__import__(module), name, None)
        return fn(*args, **kw) if callable(fn) else default
    except Exception as ex:  # one feature's fault must not stop the report (or the site build)
        print(f"Website report: {module}.{name} failed ({type(ex).__name__}: {ex}); its section is left short.")
        return default


def read_json(path):
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def nice(iso):
    """'24 Oct 2026' from an ISO date or time, 'Mar 2027' when only the month is known ("" when it is neither)."""
    text = str(iso or "")
    try:
        if len(text) == 7:
            return dt.date.fromisoformat(text + "-01").strftime("%b %Y")
        return dt.date.fromisoformat(text[:10]).strftime("%d %b %Y").lstrip("0")
    except ValueError:
        return ""


def note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic, run.font.size, run.font.color.rgb = True, Pt(9), rgb(MUTED)
    return p


def coverage_section(doc, cfg, windows):
    """World map: located stories per ISSMGE region in the latest window, and member societies without any."""
    if not windows:
        doc.add_paragraph("No stories have been archived yet.")
        return
    w = windows[0]
    rows = feature("geomap", "coverage", w["stories"], cfg)
    if not rows:
        doc.add_paragraph("The world map's regional figures are not available.")
        return
    name = lambda c: feature("geomap", "country_name", c, default=c) or c
    total = len(w["stories"])
    located = sum(r["stories"] for r in rows)
    global_ = sum(1 for s in w["stories"] if "country" in s and not s.get("country"))
    waiting = sum(1 for s in w["stories"] if "country" not in s)
    members, covered = sum(r["members"] for r in rows), sum(r["members_covered"] for r in rows)
    doc.add_paragraph(f"{w['win']['label']} ({w['win']['span']}): {located} of {total} stories are located in a country "
                      f"of one of ISSMGE's six regions; {global_} are global or about several countries and {waiting} have "
                      f"no location yet. {covered} of {members} member-society countries have at least one story.")
    table(doc, ["ISSMGE region", "Stories", "Member societies covered", "Member-society countries with no stories"],
          [(r["region"], r["stories"], f"{r['members_covered']} of {r['members']}",
            ", ".join(name(c) for c in r["gaps"]) or "—") for r in rows], [3.2, 1.8, 3.2, 8.8])
    note(doc, "A story's country is where its event, project or site is, as Claude read it (not the outlet's country "
              "or the Google News edition). The map page shows every country and period.")


def developing_section(doc, groups):
    """The developing stories with the most reports, the most recently active first among equals."""
    if not groups:
        doc.add_paragraph("No developing stories yet: they appear when a new report follows up an earlier story in "
                          "the archive (an inquiry into a landslide reported last month, for example).")
        return
    today = dt.date.today()
    active = sum(1 for g in groups if str(g.get("last", "")) >= (today - dt.timedelta(days=90)).isoformat())
    reports = sum(len(g.get("stories") or []) for g in groups)
    doc.add_paragraph(f"{len(groups)} developing stories link {reports} reports; {active} had a new report in the last "
                      "90 days. The most followed:")
    top = sorted(groups, key=lambda g: (len(g.get("stories") or []), str(g.get("last", ""))), reverse=True)[:10]
    rows = []
    for g in top:
        last = (g.get("stories") or [{}])[-1]
        stage = (last.get("stage") or "Update") if last.get("storyline") == g.get("id") else (last.get("stage") or "Event")
        where = ", ".join(x for x in (g.get("place"), feature("geomap", "country_name", g.get("country") or "",
                                                               default=g.get("country"))) if x)
        rows.append((g.get("title", ""), where or "—", len(g.get("stories") or []), nice(g.get("first")),
                     nice(g.get("last")), stage))
    table(doc, ["Developing story", "Where", "Reports", "First", "Latest", "Latest stage"], rows,
          [5.4, 3.4, 1.6, 2.1, 2.1, 2.4])


def events_section(doc, cfg, data_dir):
    """The events calendar: when it was refreshed, what it holds and the next deadlines and events."""
    events = feature("events", "load", cfg)
    if events is None:
        doc.add_paragraph("The events calendar is not available.")
        return
    today = dt.date.today()
    found = read_json(data_dir / "events.json")
    updated = nice(found.get("updated"))
    editors = sum(1 for e in events if e.get("source") == "editors")
    doc.add_paragraph(f"The calendar holds {len(events)} events: {editors} from the editors' list (data/events.yaml) and "
                      f"{len(events) - editors} found by the weekly refresh, "
                      + (f"last refreshed on {updated}." if updated else "which has not run yet (it needs the "
                         "ANTHROPIC_API_KEY secret and Anthropic's web search and web fetch tools)."))
    calls = feature("events", "deadlines", cfg, today, days=90, default=[]) or []
    if calls:
        heading(doc, "Deadlines in the next 90 days (calls for papers, registration)", 2)
        table(doc, ["Deadline", "Days left", "What is due", "Event"],
              [(nice(c["date"]) or c["date"], max(0, (dt.date.fromisoformat(c["date"]) - today).days)
                if len(c["date"]) == 10 else "date TBC", c["label"], c["event"].get("short") or c["event"].get("name", ""))
               for c in calls[:10]], [2.6, 2, 4.4, 8])
    else:
        doc.add_paragraph("No deadline (call for papers, registration) falls in the next 90 days.")
    coming = feature("events", "upcoming", cfg, today, days=365, default=[]) or []
    if coming:
        heading(doc, "Events in the next 12 months", 2)
        table(doc, ["Starts", "Event", "Where", "Kind"],
              [(nice(e["start"]) or e["start"], e.get("name", ""),
                ", ".join(x for x in (e.get("city"), feature("geomap", "country_name", e.get("country") or "",
                                                             default=e.get("country"))) if x)
                or ("Online" if e.get("online") else "—"), e.get("kind", "")) for e in coming[:10]],
              [2.4, 8, 3.6, 3])


def email_section(doc, cfg, data_dir):
    """Alerts and the weekly summary: what data/notify.json records as sent."""
    rec = read_json(data_dir / "notify.json")
    alerted = rec.get("alerted") if isinstance(rec.get("alerted"), dict) else {}
    weekly = rec.get("weekly") if isinstance(rec.get("weekly"), dict) else {}
    opts = feature("notify", "settings", cfg, default={}) or {}
    plan = []
    if opts.get("story_alerts", True):
        level = opts.get("min_relevance", 5)
        plan.append(f"an alert after any scan that finds a new story rated {level}/5" + (" or higher" if level < 5 else ""))
    if opts.get("weekly_summary", True):
        plan.append(f"the weekly summary every {opts.get('weekday', 'Monday')}")
    doc.add_paragraph("Settings (config.yaml, alerts): " + (" and ".join(plan) if plan else "all e-mail switched off") + ". "
                      "Whether e-mail is set up (the NOTIFY_TOKEN or SMTP secrets) cannot be seen when this report is "
                      "made; below is what the news desk recorded as sent.")
    stories = [k for k in alerted if not str(k).startswith("digest:")]
    days = sorted({str(v)[:10] for v in alerted.values()})
    rows = [("Stories alerted (last 60 days)", len(stories)),
            ("Latest alert", nice(days[-1]) if days else "none recorded"),
            ("Weekly summaries e-mailed", len(weekly))]
    table(doc, ["Item", "Value"], rows, [8, 9])
    if weekly:
        when = lambda iso: (nice(iso) + (f", {str(iso)[11:16]} UTC" if len(str(iso)) > 15 else "")) or str(iso)
        table(doc, ["Week", "E-mailed"], [(k, when(v)) for k, v in sorted(weekly.items(), reverse=True)[:8]], [5, 12])
    else:
        doc.add_paragraph("No weekly summary has been e-mailed yet. Every week's summary is still on the website "
                          "(Archive → Weekly summaries).")


def write_report(path, cfg, digests, windows, theme_color, site_url, changelog_path, groups=None):
    """groups: storylines.groups() of the archive (the developing stories), worked out once by the site build."""
    log = yaml.safe_load(changelog_path.read_text(encoding="utf-8")) if changelog_path.exists() else {}
    latest = digests[0] if digests else None
    data_dir = changelog_path.parent
    groups = groups if groups is not None else (feature("storylines", "groups", digests, default=[]) or [])
    numbers = iter(range(1, 100))

    def section_heading(text):
        return heading(doc, f"{next(numbers)}. {text}")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = section.bottom_margin = Cm(1.8)
    page_number_footer(section)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)

    # ---- cover block
    p = doc.add_paragraph()
    run = p.add_run("GEOTECHNICAL NEWS DESK")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10), rgb("9A7B1C")
    title = doc.add_paragraph()
    run = title.add_run("Website report: updates and review")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(24), rgb(NAVY)
    p = doc.add_paragraph()
    run = p.add_run(f"{(cfg.get('magazine') or {}).get('name', 'ISSMGE Magazine')} · generated "
                    f"{dt.datetime.now(dt.timezone.utc):%d %B %Y, %H:%M} UTC · updated automatically after every scan")
    run.font.size, run.font.color.rgb = Pt(9.5), rgb(MUTED)
    if site_url:
        p = doc.add_paragraph()
        run = p.add_run(f"Website: {site_url}")
        run.font.size, run.font.color.rgb = Pt(9.5), rgb("1D5A8A")
    bar = doc.add_paragraph()
    bar_pr = bar._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "1"), ("color", ACCENT)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    bar_pr.append(bdr)

    # ---- at a glance
    section_heading("At a glance")
    total = sum(len(d["stories"]) for d in digests)
    weekly = read_json(data_dir / "notify.json").get("weekly")
    rows = [("Latest scan", f"{latest['date']} ({latest.get('generated', '')[11:16]} UTC)" if latest else "none yet"),
            ("Stories in latest digest", len(latest["stories"]) if latest else 0),
            ("New items scanned", latest.get("new_items", 0) if latest else 0),
            ("Triage method", ("Claude AI" if latest.get("ai_used") else "Keywords only") if latest else "—"),
            ("Stories archived", total), ("Daily digests archived", len(digests)),
            ("Issue windows", len(windows)),
            ("Stories located in a country (world map)",
             f"{sum(1 for d in digests for s in d['stories'] if s.get('country'))} of {total}"),
            ("Developing stories (two or more reports)", len(groups)),
            ("Weekly summaries e-mailed", len(weekly) if isinstance(weekly, dict) else 0),
            ("Websites failing on latest scan", f"{len(latest.get('failures', []))} of {len(cfg.get('feeds', []))}" if latest else "—")]
    table(doc, ["Item", "Value"], rows, [8, 9])

    # ---- features
    if log.get("features"):
        heading(doc, "What the website does")
        for item in log["features"]:
            name, _, text = item.partition(":")
            bullet(doc, text.strip(), bold_prefix=name.strip() + ": ") if text else bullet(doc, item)

    # ---- review
    section_heading("Review of the latest scan")
    marks = {"ok": ("✔ ", "1E7B4F"), "warn": ("⚠ ", "B5560E"), "info": ("• ", MUTED)}
    for level, text in review_findings(latest, cfg, windows):
        p = doc.add_paragraph()
        mark, color = marks[level]
        r = p.add_run(mark)
        r.bold, r.font.color.rgb = True, rgb(color)
        p.add_run(text)
        p.paragraph_format.space_after = Pt(3)
    if latest and latest.get("notes"):
        heading(doc, "Scanner notes", 2)
        for n in latest["notes"]:
            bullet(doc, n)

    # ---- source health
    section_heading("Source health")
    if latest and latest.get("failures"):
        doc.add_paragraph("These websites could not be read on the latest scan. Their stories may still arrive "
                          "through the Google News searches.")
        table(doc, ["Source", "Problem", "Suggested action"],
              [(n, e[:90], source_advice(e)) for n, e in latest["failures"]], [4, 5.5, 7.5])
    else:
        doc.add_paragraph("All websites were read successfully on the latest scan.")

    # ---- archive
    section_heading("Archive by issue window")
    table(doc, ["Window", "Months", "Digests", "Stories", "Rated 5/5", "Feature leads"],
          [(w["win"]["label"], w["win"]["span"], len(w["digests"]), len(w["stories"]),
            sum(1 for s in w["stories"] if s.get("relevance") == 5),
            sum(1 for s in w["stories"] if s.get("use") == "Feature lead")) for w in windows] or [("—",) * 6],
          [2.5, 4.5, 2, 2, 2.5, 3])

    if windows:
        w = windows[0]
        section_heading(f"{w['win']['label']} by theme")
        counts = Counter(s["theme"] for s in w["stories"])
        themes = [t for t in list(cfg["themes"]) + ["General"] if counts.get(t)]
        table(doc, ["Theme", "Stories", "Share", "Rated 4–5"],
              [(t, counts[t], f"{counts[t] / len(w['stories']):.0%}",
                sum(1 for s in w["stories"] if s["theme"] == t and (s.get("relevance") or 0) >= 4)) for t in themes],
              [8, 2.5, 2.5, 3], first_col_colors=[theme_color(t) for t in themes])

    # ---- world map
    section_heading("Regional coverage (world map)" + (f": {windows[0]['win']['label']}" if windows else ""))
    coverage_section(doc, cfg, windows)

    # ---- developing stories
    section_heading("Developing stories")
    developing_section(doc, groups)

    # ---- research papers
    papers = [dict(pp, _digest=d["date"]) for w in windows[:1] for d in w["digests"] for pp in d.get("papers", [])]
    section_heading("Research papers" + (f" ({windows[0]['win']['label']})" if windows else ""))
    if papers:
        journals = Counter(pp.get("journal", "") for pp in papers)
        doc.add_paragraph(f"{len(papers)} new papers from {len(journals)} journals were picked up for Academia-to-Industry topics.")
        table(doc, ["Journal", "Papers"], journals.most_common(), [12, 3])
        best = sorted(papers, key=lambda pp: -(pp.get("relevance") or 0))[:8]
        for pp in best:
            bullet(doc, f" — {pp.get('journal', '')}, {pp.get('published', '')}" + (f". {pp['angle']}" if pp.get("angle") else ""),
                   bold_prefix=pp.get("title", ""))
    else:
        doc.add_paragraph("No research papers have been collected in this issue window yet.")

    # ---- issue briefs
    section_heading("Issue briefs")
    briefs = []
    for f in sorted((data_dir / "briefs").glob("*.json"), reverse=True):
        try:
            b = json.loads(f.read_text(encoding="utf-8"))
        except ValueError:
            continue
        briefs.append((f.stem, b.get("status", "draft"), str(b.get("generated", ""))[:16].replace("T", " "),
                       f"{b.get('stories_considered', b.get('n_stories', '—'))} stories"))
    if briefs:
        doc.add_paragraph("Claude drafts a brief for each issue window: a proposed line-up per theme, feature leads, "
                          "Academia-to-Industry topics, a regional-balance check, gaps worth covering, and the window's "
                          "developing stories and events.")
        table(doc, ["Issue window", "Status", "Generated (UTC)", "Drafted from"], briefs, [3.5, 2.5, 5, 6])
    else:
        doc.add_paragraph("No issue brief has been generated yet. Drafts appear automatically in the last weeks of each "
                          "quarter, a final brief right after it closes, and any time from the issue page.")

    # ---- events calendar
    section_heading("Events calendar")
    events_section(doc, cfg, data_dir)

    # ---- e-mail
    section_heading("E-mail alerts and weekly summary")
    email_section(doc, cfg, data_dir)

    # ---- change log
    section_heading("Change log")
    for entry in log.get("changes", []):
        p = doc.add_paragraph()
        r = p.add_run(f"{entry['date']} · {entry['title']}")
        r.bold, r.font.color.rgb = True, rgb("23466B")
        p.paragraph_format.space_before = Pt(8)
        for item in entry.get("items", []):
            bullet(doc, item)

    # ---- open items
    if log.get("open_items"):
        section_heading("Open items and recommendations")
        for item in log["open_items"]:
            bullet(doc, item)

    doc.save(path)
